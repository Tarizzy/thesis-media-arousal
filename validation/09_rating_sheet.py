"""09_rating_sheet.py - blind rating sheet for the human validation of the arousal measure.
Writes validation/09_rating_sheet.xlsx (send this to the coders) and validation/09_rating_key.csv
(keep this to yourself: it links item numbers to outlets, arousal scores and factuality).
Everything else is read-only."""
import os, re, html, sys
import numpy as np, pandas as pd, duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.environ.get("THESIS_ROOT", os.path.expanduser("~/Desktop/thesis"))
OUT = f"{ROOT}/validation"
PARQUET = f"{ROOT}/data/misinfo-general/data/*.parquet"
METADATA = f"{ROOT}/data/misinfo-general/metadata.db"
SEED, N_MAIN, N_PRACTICE, WORDS = 42, 120, 10, 180
MASK_RE = re.compile(r"<copyright>|<twitter>|<url>|(?<!\w)<?selfref>?(?!\w)", re.IGNORECASE)
FACT = {"Very Low": 0, "Low": 1, "Mixed": 2, "Mostly Factual": 3, "High": 4, "Very High": 4}
os.makedirs(OUT, exist_ok=True)


def stop(msg):
    print(f"STOP: {msg}")
    sys.exit(1)


t = pd.read_parquet(f"{ROOT}/pipeline/analysis_table.parquet")
need = [c for c in ("article_id", "source", "mbfc_factuality", "arousal") if c not in t.columns]
if need:
    stop(f"analysis_table.parquet has no column(s) {need}. It has: {list(t.columns)}")
t["fact"] = t.mbfc_factuality.map(FACT)
print(f"09_rating_sheet - {len(t):,} articles, {t.source.nunique()} outlets")

flagged = []
try:
    o = pd.read_csv(f"{ROOT}/models/outlet_level.csv", index_col=0)
    if "outlet_ok" in o:
        flagged = list(o.index[~o.outlet_ok.astype(bool)])
except FileNotFoundError:
    pass
pool = t[~t.source.isin(flagged)].copy()
print(f"   excluded {len(flagged)} flagged outlet(s): {flagged}")

pool["stratum"] = pd.qcut(pool.arousal, 5, labels=[1, 2, 3, 4, 5]).astype(int)
pool = pool.sample(frac=1, random_state=SEED)
used, picked, short = {}, [], []
for s in range(1, 6):
    for f in range(5):
        cell = pool[(pool.stratum == s) & (pool.fact == f)]
        got = 0
        for _, r in cell.iterrows():
            if used.get(r.source, 0) >= 2:
                continue
            used[r.source] = used.get(r.source, 0) + 1
            picked.append(r)
            got += 1
            if got == 5:
                break
        if got < 5:
            short.append(f"stratum {s} x class {f}: {got}/5")
rest = pool[~pool.article_id.isin([r.article_id for r in picked])]
for _, r in rest.iterrows():
    if len(picked) >= N_MAIN + N_PRACTICE:
        break
    if used.get(r.source, 0) < 2:
        used[r.source] = used.get(r.source, 0) + 1
        picked.append(r)
sample = pd.DataFrame(picked)
if len(sample) < N_MAIN + N_PRACTICE:
    stop(f"only {len(sample)} articles could be sampled, need {N_MAIN + N_PRACTICE}")
if short:
    print("   info  thin cells topped up from the pool: " + "; ".join(short))

con = duckdb.connect()
con.register("want", sample[["article_id", "source"]])
txt = con.execute(f"""SELECT p.article_id, p.title, p.content FROM read_parquet('{PARQUET}') p
                      JOIN want w ON w.article_id = p.article_id""").df()
sample = sample.merge(txt, on="article_id", how="inner")
if len(sample) < N_MAIN + N_PRACTICE:
    stop(f"only {len(sample)} of the sampled articles have text in the corpus parquet")


def blind(text, source):
    text = MASK_RE.sub(" ", html.unescape(str(text)))
    text = re.sub(r"https?://\S+|www\.\S+", " ", text)
    if len(source) >= 5:
        pat = r"\b" + r"[\s\-\.]?".join(re.escape(c) for c in source) + r"\b"
        text = re.sub(pat, "[outlet]", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    words = text.split(" ")
    if len(words) > WORDS:
        text = " ".join(words[:WORDS])
        cut = max(text.rfind(". "), text.rfind("! "), text.rfind("? "))
        text = (text[:cut + 1] if cut > len(text) * 0.6 else text) + " [...]"
    return text


sample = sample.sample(frac=1, random_state=SEED + 1).reset_index(drop=True)
sample["excerpt"] = [blind(c, s) for c, s in zip(sample.content, sample.source)]
sample["headline"] = [blind(h, s) for h, s in zip(sample.title, sample.source)]
practice = sample.iloc[:N_PRACTICE].copy()
main = sample.iloc[N_PRACTICE:N_PRACTICE + N_MAIN].copy()
practice["item"] = [f"P{i:02d}" for i in range(1, len(practice) + 1)]
main["item"] = [f"R{i:03d}" for i in range(1, len(main) + 1)]

cl = pd.read_csv(f"{ROOT}/partisanship/cluster_framing.csv", index_col=0)
wc = pd.read_parquet(f"{ROOT}/partisanship/word_counts.parquet")
tot = wc.groupby("word").n.sum()
glob = tot / tot.sum()
words_by_cluster = {}
for c, g in wc.groupby("cluster"):
    s = g.groupby("word").n.sum()
    s = s[s >= 200]
    if len(s):
        lift = (s / s.sum()) / glob.reindex(s.index)
        words_by_cluster[int(c)] = ", ".join(lift.sort_values(ascending=False).index[:12])
cols_corpus = con.execute(f"DESCRIBE SELECT * FROM read_parquet('{PARQUET}') LIMIT 1").df().column_name.tolist()
if "topic_id" in cols_corpus:
    heads = con.execute(f"""
        WITH a AS (SELECT article_id, title, topic_id FROM read_parquet('{PARQUET}')
                   WHERE title IS NOT NULL AND length(title) > 25)
        SELECT tc.cluster, a.title,
               row_number() OVER (PARTITION BY tc.cluster ORDER BY hash(a.article_id)) AS rn
        FROM a JOIN '{ROOT}/clustering/topic_clusters.parquet' tc ON tc.topic_id = a.topic_id""").df()
else:
    con.execute(f"ATTACH '{METADATA}' AS meta (READ_ONLY)")
    heads = con.execute(f"""
        SELECT tc.cluster, m.title,
               row_number() OVER (PARTITION BY tc.cluster ORDER BY hash(m.article_id)) AS rn
        FROM meta.articles m JOIN '{ROOT}/clustering/topic_clusters.parquet' tc ON tc.topic_id = m.topic_id
        WHERE m.title IS NOT NULL AND length(m.title) > 25""").df()
ex = (heads[heads.rn <= 3].groupby("cluster").title
      .apply(lambda s: "  |  ".join(html.unescape(str(x)) for x in s)).to_dict())
topics = pd.DataFrame({"topic": cl.index,
                       "name": [str(cl.loc[c, "label"]) if "label" in cl else "" for c in cl.index],
                       "characteristic words": [words_by_cluster.get(int(c), "") for c in cl.index],
                       "example headlines": [ex.get(int(c), "") for c in cl.index]})

wb = Workbook()
wrap = Alignment(wrap_text=True, vertical="top")


def sheet(ws, rows, cols, widths, note, rating_col):
    ws["A1"] = note
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(cols)
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = "A4"
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{rating_col}4:{rating_col}{3 + len(rows)}")


ws = wb.active
ws.title = "Instructions"
for line in ["Rating news writing for emotional language", "",
             "Rate THE LANGUAGE of each excerpt, 0 to 2, in the 'rating' column.",
             "  0 = neutral reporting: plain statement, claims attributed, few intensifiers.",
             "  1 = some emotive language: scattered charged words or evaluative framing.",
             "  2 = strongly emotive: sustained alarm, outrage, mockery or moral condemnation;",
             "      exclamation marks, capitals, rhetorical questions, appeals to the reader.", "",
             "Rate the wording, not the subject: a dramatic event described calmly is low.",
             "Rate the wording, not whether you agree with it or believe it.",
             "Include the headline and any quoted speech.",
             "Rate only the text shown, even where it is cut off. Please do not look up the original.",
             "Work alone after the practice items. First instinct; do not revisit or balance the totals.", "",
             "Do the 'Practice' tab first (10 items, not counted), then 'Articles' (120 items).",
             "'Topics' is optional, 15 minutes: rate each topic 0-2 for how politically contested",
             "the ISSUE was in the US between 2017 and 2022 (0 = little partisan disagreement,",
             "1 = some, 2 = cleanly divides the parties). Judge the issue, not the wording.", "",
             "Send the file back with the rating columns filled in."]:
    ws.append([line])
ws.column_dimensions["A"].width = 100

for name, df, note in [("Practice", practice, "PRACTICE - 10 items, not counted. Discuss with the other coder afterwards."),
                       ("Articles", main, "ARTICLES - 120 items. Work alone. Rate the language, 0 to 2.")]:
    sheet(wb.create_sheet(name), [[r.item, r.headline, r.excerpt, None] for r in df.itertuples()],
          ["item", "headline", "excerpt", "rating"], [8, 45, 110, 10], note, "D")
sheet(wb.create_sheet("Topics"), topics.values.tolist(),
      ["topic", "name", "characteristic words", "example headlines", "rating"], [8, 26, 60, 70, 10],
      "TOPICS (optional) - rate how politically contested each ISSUE was, 0 to 2.", "E")
wb.save(f"{OUT}/09_rating_sheet.xlsx")

key = pd.concat([practice.assign(set="practice"), main.assign(set="main")])
key[["item", "set", "article_id", "source", "mbfc_factuality", "fact", "arousal", "stratum"]] \
    .to_csv(f"{OUT}/09_rating_key.csv", index=False)
print(f"   {len(main)} rating items + {len(practice)} practice, {sample.source.nunique()} outlets, "
      f"{len(topics)} topics")
print(f"   arousal range covered: {sample.arousal.min():.3f} to {sample.arousal.max():.3f}")
print("RESULT: wrote validation/09_rating_sheet.xlsx (send) and 09_rating_key.csv (keep private)")
